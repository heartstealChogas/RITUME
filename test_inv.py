import openpyxl
import sys

def debug_inv(path):
    print(f"Reading: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    
    print("First 15 rows:")
    for r in range(1, min(ws.max_row, 15) + 1):
        row_vals = [str(ws.cell(row=r, column=c).value or '').replace(' ', '') for c in range(1, min(ws.max_column, 20) + 1)]
        print(f"R{r}: {row_vals}")

if __name__ == '__main__':
    debug_inv(r"C:\Users\admin\Desktop\rie\data\송장\영도점\영도점_송장_20251001.xlsx")
