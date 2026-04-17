import openpyxl
from pathlib import Path
import json

def dump_data():
    po_path = Path(r"c:\Users\admin\Desktop\rie\data\발주서\영도점\영도점_발주서_20251001.xlsx")
    inv_path = Path(r"c:\Users\admin\Desktop\rie\data\송장\영도점\영도점_송장_20251001.xlsx")
    
    # We might have named them differently. Let's find it.
    po_dir = Path(r"c:\Users\admin\Desktop\rie\data\발주서")
    inv_dir = Path(r"c:\Users\admin\Desktop\rie\data\송장")
    
    out = {}
    for inv_f in inv_dir.rglob("*.xlsx"):
        if not inv_f.name.startswith("~$") and "20251001" in inv_f.name and "영도점" in str(inv_f):
            wb = openpyxl.load_workbook(inv_f, data_only=True)
            ws = wb.active
            rows = []
            for r in range(1, min(ws.max_row, 15) + 1):
                rows.append([str(ws.cell(row=r, column=c).value or '') for c in range(1, 20)])
            out['inv'] = rows
            
    for po_f in po_dir.rglob("*.xlsx"):
        if not po_f.name.startswith("~$") and "20251001" in po_f.name and "영도점" in str(po_f):
            wb = openpyxl.load_workbook(po_f, data_only=True)
            ws = wb.active
            rows = []
            for r in range(1, min(ws.max_row, 5) + 1):
                rows.append([str(ws.cell(row=r, column=c).value or '') for c in range(1, 20)])
            out['po'] = rows

    with open(r"c:\Users\admin\Desktop\rie\debug_out.json", "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

if __name__ == '__main__':
    dump_data()
