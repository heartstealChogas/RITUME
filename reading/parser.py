import re
from pathlib import Path
import openpyxl

def parse_purchase_order(path: Path) -> list[dict]:
    """
    Parses a single Excel Purchase Order and returns extracted products and metadata.
    Automatically handles formatting shifts by searching for known headers.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    contact_person = ''
    contact_tel = ''
    supplier_address = ''

    # First Pass: Find Metadata (Contact, TEL, Address)
    for r in range(1, min(ws.max_row, 50) + 1):
        for c in range(1, 10):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if not val: continue
            
            # Finding TEL
            if 'TEL' in val.upper() and not contact_tel:
                contact_tel = str(ws.cell(row=r, column=c+1).value or '').strip()
                
            # Finding 담당자
            if val in ('담당자', '담당자명'):
                contact_person = str(ws.cell(row=r, column=c+1).value or '').strip()
                
            # Finding 납품처 주소
            if '납품처 주소' in val or '주소' in val:
                if not supplier_address:
                    supplier_address = str(ws.cell(row=r, column=c+1).value or '').strip()

    # Second Pass: Find "No." marker for Product rows
    start_row = 0
    name_col = 3    # Default C
    barcode_col = 4 # Default D
    qty_col = 7     # Default G
    
    for r in range(1, ws.max_row + 1):
        val_a = str(ws.cell(row=r, column=1).value or '').strip().upper()
        if val_a in ('NO.', 'NO'):
            start_row = r + 1
            # Discover exact columns dynamically
            for c_idx in range(1, ws.max_column + 1):
                hdr_val = str(ws.cell(row=r, column=c_idx).value or '').replace('\n', '').replace(' ', '').strip()
                if '제품명' in hdr_val or '상품명' in hdr_val:
                    name_col = c_idx
                elif '바코드' in hdr_val:
                    barcode_col = c_idx
                elif '총' in hdr_val and '주문수량' in hdr_val:
                    qty_col = c_idx
            break

    if start_row == 0:
        raise ValueError("엑셀 내에서 제품 목록(No.) 테이블을 찾을 수 없습니다.")

    # Third Pass: Extract Products
    extracted = []
    for r in range(start_row, ws.max_row + 1):
        val_a = str(ws.cell(row=r, column=1).value or '').strip().upper()
        if val_a == 'TOTAL':
            break # End of table
        
        prod_name = str(ws.cell(row=r, column=name_col).value or '').strip()
        raw_qty = str(ws.cell(row=r, column=qty_col).value or '')
        
        # safely parse numbers like "200ea"
        digits = re.sub(r'[^\d]', '', raw_qty)
        qty = int(digits) if digits else 0
            
        if prod_name and qty > 0:
            barcode_raw = ws.cell(row=r, column=barcode_col).value
            if barcode_raw is None:
                barcode = ''
            elif isinstance(barcode_raw, float) and barcode_raw.is_integer():
                barcode = str(int(barcode_raw))
            else:
                barcode = str(barcode_raw).strip()
            extracted.append({
                'name': prod_name,
                'barcode': barcode,
                'qty': qty
            })

    # Prepare standard export row items
    formatted_rows = []
    seq_gir = 1
    
    for idx, item in enumerate(extracted):
        out_row = {
            '순번': idx + 1,
            '거래처코드': 'C0000010',
            '출하의뢰번호 (NO_GIR)': 1,
            '출하의뢰항번 (SEQ_GIR)': seq_gir,
            '납품처코드': 'C0000010',
            '상품명(사이트)': '',
            '사이트명': '',
            '상품명 (NM_ITEM)_1': item['name'],
            '바코드 (CD_ITEM)_1': item['barcode'],
            '주문수량 (QT_GIR)': item['qty'],
            '주문자명 (NM_CUST)': contact_person,
            '주문자 연락처1 (NO_TEL_D1)': contact_tel,
            '수취인명 (NM_CUST_DLV)': contact_person,
            '수취인연락처1 (NO_TEL_D1)': contact_tel,
            '배송주소': supplier_address,
        }
        formatted_rows.append(out_row)
        seq_gir += 1
        
    return formatted_rows
