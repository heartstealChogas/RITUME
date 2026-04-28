import re
from pathlib import Path
import openpyxl

OUT_HEADERS = [
    '순번', '거래처코드', '출하의뢰번호 (NO_GIR)', '출하의뢰항번 (SEQ_GIR)',
    '납품처코드', '상품명(사이트)', '사이트명', '상품명 (NM_ITEM)_1',
    '바코드 (CD_ITEM)_1', '상품명 (NM_ITEM)_2', '바코드 (CD_ITEM)_2',
    '상품명 (NM_ITEM)_3', '바코드 (CD_ITEM)_3', '주문수량 (QT_GIR)',
    '주문자명 (NM_CUST)', '주문자 연락처1 (NO_TEL_D1)', '주문자 연락처2 (NO_TEL_D2)',
    '수취인명 (NM_CUST_DLV)', '수취인연락처1 (NO_TEL_D1)', '수취인연락처2 (NO_TEL_D2)',
    '수취인우편번호 CD_ZIP', '배송주소', '수취인주소2 (ADDR2)', '배송메시지 (DC_REQ)',
    '송장번호 (CD_INVOICE)', '택배사명 (CD_DELI)', '주문번호', '반품사유'
]


def parse_purchase_order(path: Path) -> list[dict]:
    """
    발주서 엑셀을 파싱하여 송장 양식 행 목록을 반환.
    헤더 위치를 동적으로 탐지하므로 서식 변형에 대응 가능.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    contact_person = ''
    contact_tel = ''
    supplier_address = ''

    # 1단계: 담당자 / TEL / 납품처 주소 추출
    for r in range(1, min(ws.max_row, 50) + 1):
        for c in range(1, 10):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if not val:
                continue
            if 'TEL' in val.upper() and not contact_tel:
                contact_tel = str(ws.cell(row=r, column=c + 1).value or '').strip()
            if val in ('담당자', '담당자명'):
                contact_person = str(ws.cell(row=r, column=c + 1).value or '').strip()
            if ('납품처 주소' in val or '주소' in val) and not supplier_address:
                supplier_address = str(ws.cell(row=r, column=c + 1).value or '').strip()

    # 2단계: No. 헤더 행 탐지 → 제품 테이블 시작
    start_row = 0
    name_col = 3
    barcode_col = 4
    qty_col = 7

    for r in range(1, ws.max_row + 1):
        val_a = str(ws.cell(row=r, column=1).value or '').strip().upper()
        if val_a in ('NO.', 'NO'):
            start_row = r + 1
            for c_idx in range(1, ws.max_column + 1):
                hdr = str(ws.cell(row=r, column=c_idx).value or '').replace('\n', '').replace(' ', '').strip()
                if '제품명' in hdr or '상품명' in hdr:
                    name_col = c_idx
                elif '바코드' in hdr:
                    barcode_col = c_idx
                elif '총' in hdr and '주문수량' in hdr:
                    qty_col = c_idx
            break

    if start_row == 0:
        raise ValueError("엑셀에서 제품 목록(No.) 테이블을 찾을 수 없습니다.")

    # 3단계: 제품 행 추출
    extracted = []
    for r in range(start_row, ws.max_row + 1):
        val_a = str(ws.cell(row=r, column=1).value or '').strip().upper()
        if val_a == 'TOTAL':
            break
        prod_name = str(ws.cell(row=r, column=name_col).value or '').strip()
        raw_qty = str(ws.cell(row=r, column=qty_col).value or '')
        digits = re.sub(r'[^\d]', '', raw_qty)
        qty = int(digits) if digits else 0
        if prod_name and qty > 0:
            barcode = str(ws.cell(row=r, column=barcode_col).value or '').strip()
            extracted.append({'name': prod_name, 'barcode': barcode, 'qty': qty})

    # 4단계: 송장 양식 행 생성
    formatted_rows = []
    for idx, item in enumerate(extracted):
        formatted_rows.append({
            '순번': idx + 1,
            '거래처코드': 'C0000010',
            '출하의뢰번호 (NO_GIR)': 1,
            '출하의뢰항번 (SEQ_GIR)': idx + 1,
            '납품처코드': 'C0000010',
            '상품명(사이트)': '',
            '사이트명': '',
            '상품명 (NM_ITEM)_1': item['name'],
            '바코드 (CD_ITEM)_1': item['barcode'],
            '상품명 (NM_ITEM)_2': '',
            '바코드 (CD_ITEM)_2': '',
            '상품명 (NM_ITEM)_3': '',
            '바코드 (CD_ITEM)_3': '',
            '주문수량 (QT_GIR)': item['qty'],
            '주문자명 (NM_CUST)': contact_person,
            '주문자 연락처1 (NO_TEL_D1)': contact_tel,
            '주문자 연락처2 (NO_TEL_D2)': '',
            '수취인명 (NM_CUST_DLV)': contact_person,
            '수취인연락처1 (NO_TEL_D1)': contact_tel,
            '수취인연락처2 (NO_TEL_D2)': '',
            '수취인우편번호 CD_ZIP': '',
            '배송주소': supplier_address,
            '수취인주소2 (ADDR2)': '',
            '배송메시지 (DC_REQ)': '',
            '송장번호 (CD_INVOICE)': '',
            '택배사명 (CD_DELI)': '',
            '주문번호': '',
            '반품사유': '',
        })

    return formatted_rows
