import sys, os
from pathlib import Path
import openpyxl

def debug_inv():
    # Use proper paths
    inv_dir = Path(r"c:\Users\admin\Desktop\rie\data\송장")
    for inv_f in inv_dir.rglob("*.xlsx"):
        if not inv_f.name.startswith("~$"):
            print(f"Reading: {inv_f}")
            wb = openpyxl.load_workbook(inv_f, data_only=True)
            ws = wb.active
            
            print("First 15 rows:")
            for r in range(1, min(ws.max_row, 15) + 1):
                row_vals = [str(ws.cell(row=r, column=c).value or '').replace(' ', '') for c in range(1, min(ws.max_column, 20) + 1)]
                print(f"R{r}: {row_vals}")
            break

if __name__ == '__main__':
    debug_inv()
