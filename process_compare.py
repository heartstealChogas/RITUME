import os
import re
import json
from pathlib import Path
import openpyxl
import psycopg2

from reading.parser import parse_purchase_order

BASE_DIR = Path(r"C:\Users\admin\Desktop\rie")
DATA_DIR = BASE_DIR / "data"

def setup_db():
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL environment variable is not set")
    conn = psycopg2.connect(database_url)
    cur = conn.cursor()
    cur.execute('''
        CREATE TABLE IF NOT EXISTS comparisons (
            id SERIAL PRIMARY KEY,
            branch_name TEXT,
            file_date TEXT,
            status TEXT,
            seq TEXT,
            name TEXT,
            barcode TEXT,
            po_qty INTEGER,
            inv_no TEXT,
            inv_qty INTEGER,
            qty_diff INTEGER
        )
    ''')
    conn.commit()
    cur.close()
    return conn

def parse_invoice(path: Path) -> list[dict]:
    """
    Parses an ERP-format invoice Excel file.
    The invoice files use ERP column headers:
      NM_ITEM_1 → product name, CD_ITEM_1 → barcode,
      QT_GIR    → quantity,   CD_INVOICE → invoice number
    Header row is auto-detected within the first 15 rows.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active

    start_row = 0
    inv_col, name_col, barcode_col, qty_col, no_gir_col = -1, -1, -1, -1, -1

    # Scan first 15 rows for the ERP-format header row.
    # The header row contains identifiers like NM_ITEM, CD_ITEM, QT_GIR, CD_INVOICE.
    for r in range(1, min(ws.max_row, 15) + 1):
        row_vals = [
            re.sub(r'\s+', '', str(ws.cell(row=r, column=c).value or ''))
            for c in range(1, ws.max_column + 1)
        ]
        # Detect ERP header row by looking for ERP-style identifiers
        has_erp_header = any(
            'NM_ITEM' in h or 'CD_ITEM' in h or 'QT_GIR' in h or 'CD_INVOICE' in h
            for h in row_vals
        )
        # Also detect traditional Korean invoice headers as a fallback
        has_ko_header = any(
            '송장번호' in h or '상품명' in h or '제품명' in h or '바코드' in h
            for h in row_vals
        )
        if has_erp_header or has_ko_header:
            start_row = r + 1
            for c_idx, h in enumerate(row_vals, 1):
                # ERP-format column detection (primary)
                if 'CD_INVOICE' in h:
                    inv_col = c_idx
                elif 'NO_GIR' in h and no_gir_col == -1:
                    no_gir_col = c_idx
                elif 'NM_ITEM' in h and '_1' in h and name_col == -1:
                    name_col = c_idx
                elif 'CD_ITEM' in h and '_1' in h and barcode_col == -1:
                    barcode_col = c_idx
                elif 'QT_GIR' in h and qty_col == -1:
                    qty_col = c_idx
                # Korean fallback column detection
                elif '송장번호' in h and inv_col == -1:
                    inv_col = c_idx
                elif ('상품명' in h or '제품명' in h) and name_col == -1:
                    name_col = c_idx
                elif '바코드' in h and barcode_col == -1:
                    barcode_col = c_idx
                elif '수량' in h and qty_col == -1:
                    qty_col = c_idx
            break

    rows = []
    if start_row > 0:
        for r in range(start_row, ws.max_row + 1):
            inv_no = str(ws.cell(row=r, column=inv_col).value or '').strip() if inv_col > 0 else ''
            # Fall back to NO_GIR (delivery request number) when CD_INVOICE is empty
            if not inv_no and no_gir_col > 0:
                raw_no_gir = ws.cell(row=r, column=no_gir_col).value
                if raw_no_gir is not None:
                    val = raw_no_gir
                    if isinstance(val, float) and val.is_integer():
                        val = int(val)
                    inv_no = str(val).strip()
            name   = str(ws.cell(row=r, column=name_col).value or '').strip() if name_col > 0 else ''
            barcode_raw = ws.cell(row=r, column=barcode_col).value if barcode_col > 0 else None
            if barcode_raw is None:
                barcode = ''
            elif isinstance(barcode_raw, float) and barcode_raw.is_integer():
                barcode = str(int(barcode_raw))
            else:
                barcode = str(barcode_raw).strip()
            raw_qty = str(ws.cell(row=r, column=qty_col).value or '') if qty_col > 0 else ''
            digits = re.sub(r'[^\d]', '', raw_qty)
            qty = int(digits) if digits else 0

            if name or barcode:
                rows.append({
                    'invNo': inv_no,
                    'name': name,
                    'barcode': barcode,
                    'qty': qty
                })
    return rows

def norm_str(s):
    return re.sub(r'\s+', '', str(s)).lower()

def build_comparison(po_rows, inv_rows):
    result = []
    inv_used = set()
    
    for po in po_rows:
        try:
            seq = int(po.get('순번', 0))
        except:
            seq = po.get('순번', '')
            
        name = str(po.get('상품명 (NM_ITEM)_1', ''))
        barcode = str(po.get('바코드 (CD_ITEM)_1', ''))
        po_qty = int(po.get('주문수량 (QT_GIR)', 0))
        
        match_idx = -1
        # 1. match by barcode
        if barcode:
            for i, ir in enumerate(inv_rows):
                if i not in inv_used and ir['barcode'] == barcode:
                    match_idx = i
                    break
                    
        # 2. match by normalized name
        if match_idx == -1:
            norm_name = norm_str(name)
            for i, ir in enumerate(inv_rows):
                if i not in inv_used and norm_str(ir['name']) == norm_name:
                    match_idx = i
                    break
                    
        if match_idx != -1:
            matched_inv = inv_rows[match_idx]
            inv_used.add(match_idx)
            inv_qty = matched_inv['qty']
            qty_diff = inv_qty - po_qty
            result.append({
                'status': 'ok',
                'seq': str(seq),
                'name': name,
                'barcode': barcode,
                'po_qty': po_qty,
                'inv_no': matched_inv['invNo'],
                'inv_qty': inv_qty,
                'qty_diff': qty_diff
            })
        else:
            result.append({
                'status': 'miss',
                'seq': str(seq),
                'name': name,
                'barcode': barcode,
                'po_qty': po_qty,
                'inv_no': '',
                'inv_qty': 0,
                'qty_diff': -po_qty
            })
            
    for i, ir in enumerate(inv_rows):
        if i not in inv_used and (ir['name'] or ir['barcode']):
            result.append({
                'status': 'extra',
                'seq': '—',
                'name': ir['name'] or '—',
                'barcode': ir['barcode'] or '',
                'po_qty': 0,
                'inv_no': ir['invNo'],
                'inv_qty': ir['qty'],
                'qty_diff': ir['qty']
            })
            
    return result

def main():
    conn = setup_db()
    cur = conn.cursor()
    
    po_dir = DATA_DIR / "발주서"
    inv_dir = DATA_DIR / "송장"
    
    if not po_dir.exists() or not inv_dir.exists():
        print("Missing '발주서' or '송장' folder in data.")
        return

    po_files = list(po_dir.rglob("*.xlsx"))
    total_inserted = 0
    
    print(f"Found {len(po_files)} total PO files to scan.")
    for pf in po_files:
        if pf.name.startswith("~$"): continue
        branch_name = pf.parent.name
        date_match = re.search(r'20\d{6}', pf.name)
        if not date_match: continue
        file_date = date_match.group(0)
        
        inv_branch_dir = inv_dir / branch_name
        if not inv_branch_dir.exists():
            continue
        
        matched_inv = None
        for inv_f in inv_branch_dir.rglob("*.xlsx"):
            if not inv_f.name.startswith("~$") and file_date in inv_f.name:
                matched_inv = inv_f
                break
                
        if not matched_inv:
            print(f"Skipping {pf.name} (No matching invoice found in {branch_name})")
            continue
            
        # Check for duplicate entries
        cur.execute("SELECT COUNT(*) FROM comparisons WHERE branch_name = %s AND file_date = %s", (branch_name, file_date))
        if cur.fetchone()[0] > 0:
            print(f"Notification: duplicate value is already registered for {branch_name} - {file_date}. Skipping.")
            continue

        print(f"Comparing {branch_name} - {file_date}")
        try:
            po_data = parse_purchase_order(pf)
            inv_data = parse_invoice(matched_inv)
            
            comp_res = build_comparison(po_data, inv_data)
            
            for res in comp_res:
                cur.execute('''
                    INSERT INTO comparisons 
                    (branch_name, file_date, status, seq, name, barcode, po_qty, inv_no, inv_qty, qty_diff)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ''', (
                    branch_name, file_date, res['status'], res['seq'], res['name'],
                    res['barcode'], res['po_qty'], res['inv_no'], res['inv_qty'], res['qty_diff']
                ))
            total_inserted += len(comp_res)
        except Exception as e:
            print(f"Error processing {branch_name} {file_date}: {e}")
            
    conn.commit()

    # Generate JSON for web UI
    print("Exporting database to JSON for web UI...")
    cur.execute("SELECT * FROM comparisons")
    # Get column names
    col_names = [description[0] for description in cur.description]
    all_rows = [dict(zip(col_names, row)) for row in cur.fetchall()]
    
    json_out_path = BASE_DIR / "web" / "comparisons_data.json"
    json_out_path.parent.mkdir(exist_ok=True)
    with open(json_out_path, 'w', encoding='utf-8') as f:
        json.dump(all_rows, f, ensure_ascii=False)

    conn.close()
    print(f"Done. Inserted {total_inserted} comparison records into Supabase PostgreSQL")
    print(f"JSON data written to {json_out_path.name}")

if __name__ == "__main__":
    main()
