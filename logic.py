import re
import openpyxl
from pathlib import Path
from reading.parser import parse_purchase_order

def parse_invoice(ws) -> list[dict]:
    """
    Parses an ERP-format invoice worksheet.
    Header row is auto-detected within the first 15 rows.
    """
    start_row = 0
    inv_col, name_col, barcode_col, qty_col, no_gir_col = -1, -1, -1, -1, -1

    # Scan first 15 rows for header row
    for r in range(1, min(ws.max_row, 15) + 1):
        row_vals = [
            re.sub(r'\s+', '', str(ws.cell(row=r, column=c).value or ''))
            for c in range(1, ws.max_column + 1)
        ]
        has_erp_header = any(
            'NM_ITEM' in h or 'CD_ITEM' in h or 'QT_GIR' in h or 'CD_INVOICE' in h
            for h in row_vals
        )
        has_ko_header = any(
            '송장번호' in h or '상품명' in h or '제품명' in h or '바코드' in h
            for h in row_vals
        )
        if has_erp_header or has_ko_header:
            start_row = r + 1
            for c_idx, h in enumerate(row_vals, 1):
                if 'CD_INVOICE' in h: inv_col = c_idx
                elif 'NO_GIR' in h and no_gir_col == -1: no_gir_col = c_idx
                elif 'NM_ITEM' in h and '_1' in h and name_col == -1: name_col = c_idx
                elif 'CD_ITEM' in h and '_1' in h and barcode_col == -1: barcode_col = c_idx
                elif 'QT_GIR' in h and qty_col == -1: qty_col = c_idx
                elif '송장번호' in h and inv_col == -1: inv_col = c_idx
                elif ('상품명' in h or '제품명' in h) and name_col == -1: name_col = c_idx
                elif '바코드' in h and barcode_col == -1: barcode_col = c_idx
                elif '수량' in h and qty_col == -1: qty_col = c_idx
            break

    rows = []
    if start_row > 0:
        for r in range(start_row, ws.max_row + 1):
            inv_no = str(ws.cell(row=r, column=inv_col).value or '').strip() if inv_col > 0 else ''
            if not inv_no and no_gir_col > 0:
                val = ws.cell(row=r, column=no_gir_col).value
                if val is not None:
                    if isinstance(val, float) and val.is_integer(): val = int(val)
                    inv_no = str(val).strip()
            name = str(ws.cell(row=r, column=name_col).value or '').strip() if name_col > 0 else ''
            barcode_raw = ws.cell(row=r, column=barcode_col).value if barcode_col > 0 else None
            if barcode_raw is None: barcode = ''
            elif isinstance(barcode_raw, float) and barcode_raw.is_integer(): barcode = str(int(barcode_raw))
            else: barcode = str(barcode_raw).strip()
            raw_qty = str(ws.cell(row=r, column=qty_col).value or '') if qty_col > 0 else ''
            digits = re.sub(r'[^\d]', '', raw_qty)
            qty = int(digits) if digits else 0

            if name or barcode:
                rows.append({'invNo': inv_no, 'name': name, 'barcode': barcode, 'qty': qty})
    return rows

def norm_str(s):
    return re.sub(r'\s+', '', str(s)).lower()

def build_comparison(po_rows, inv_rows):
    result = []
    inv_used = set()
    
    for po in po_rows:
        try: seq = int(po.get('순번', 0))
        except: seq = po.get('순번', '')
            
        name = str(po.get('상품명 (NM_ITEM)_1', ''))
        barcode = str(po.get('바코드 (CD_ITEM)_1', ''))
        po_qty = int(po.get('주문수량 (QT_GIR)', 0))
        
        match_idx = -1
        if barcode:
            for i, ir in enumerate(inv_rows):
                if i not in inv_used and ir['barcode'] == barcode:
                    match_idx = i; break
        if match_idx == -1:
            norm_name = norm_str(name)
            for i, ir in enumerate(inv_rows):
                if i not in inv_used and norm_str(ir['name']) == norm_name:
                    match_idx = i; break
                    
        if match_idx != -1:
            matched_inv = inv_rows[match_idx]
            inv_used.add(match_idx)
            inv_qty = matched_inv['qty']
            result.append({
                'status': 'ok', 'seq': str(seq), 'name': name, 'barcode': barcode,
                'po_qty': po_qty, 'inv_no': matched_inv['invNo'], 'inv_qty': inv_qty, 'qty_diff': inv_qty - po_qty
            })
        else:
            result.append({
                'status': 'miss', 'seq': str(seq), 'name': name, 'barcode': barcode,
                'po_qty': po_qty, 'inv_no': '', 'inv_qty': 0, 'qty_diff': -po_qty
            })
            
    for i, ir in enumerate(inv_rows):
        if i not in inv_used and (ir['name'] or ir['barcode']):
            result.append({
                'status': 'extra', 'seq': '—', 'name': ir['name'] or '—', 'barcode': ir['barcode'] or '',
                'po_qty': 0, 'inv_no': ir['invNo'], 'inv_qty': ir['qty'], 'qty_diff': ir['qty']
            })
            
    return result
